from __future__ import absolute_import, unicode_literals

from datetime import date
from decimal import Decimal
from io import BytesIO
from openpyxl import Workbook


class XLSDocument(object):
    def __init__(self):
        self.workbook = Workbook(optimized_write=True)
        self.sheet = None

    def add_sheet(self, title=None):
        self.sheet = self.workbook.create_sheet(title=title)

    def table(self, titles, rows):
        if titles:
            self.sheet.append(titles)

        for row in rows:
            processed = []
            for i, value in enumerate(row):
                if isinstance(value, date):
                    processed.append(value.strftime('%Y-%m-%d'))
                elif isinstance(value, (int, float, Decimal)):
                    processed.append(value)
                elif value is None:
                    processed.append('-')
                else:
                    processed.append(('%s' % value).strip())

            self.sheet.append(processed)

    def to_response(self, filename):
        from django.http import HttpResponse
        output = BytesIO()
        self.workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'),
            )
        output.close()
        response['Content-Disposition'] = 'attachment; filename="%s"' % (
            filename,
            )
        return response
